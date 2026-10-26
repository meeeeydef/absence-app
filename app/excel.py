"""Excel xlsx and XML Spreadsheet 2003 (.xls) with bounded parsing."""
import io
import zipfile
from defusedxml import ElementTree as ET
from openpyxl import load_workbook

NS='urn:schemas-microsoft-com:office:spreadsheet'
MAX_ROWS=10010
MAX_COLS=100

def parse_excel(raw,filename):
    total=0;sheets=[]
    if filename.lower().endswith('.xls'):
        # The supplied internal export is XML Spreadsheet 2003, not binary XLS.
        root=ET.fromstring(raw)
        if root.tag!='{'+NS+'}Workbook': raise ValueError('XML Excel attendu')
        for sheet in root.findall('{'+NS+'}Worksheet'):
            rows=[]
            for row in sheet.findall('.//{'+NS+'}Row'):
                total+=1
                if total>MAX_ROWS: raise ValueError('Trop de lignes')
                cells=[]
                for cell in row.findall('{'+NS+'}Cell'):
                    idx=int(cell.get('{'+NS+'}Index',len(cells)+1))-1
                    if idx<len(cells) or idx>=MAX_COLS: raise ValueError('Index de colonne invalide')
                    cells.extend(['']*(idx-len(cells)))
                    data=cell.find('{'+NS+'}Data')
                    cells.append(''.join(data.itertext()).strip() if data is not None else '')
                rows.append(cells)
            if rows: sheets.append({'name':sheet.get('{'+NS+'}Name','Export'),'headers':rows[0],'rows':rows[1:]})
    else:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>30*1024*1024: raise ValueError('Archive volumineuse')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            for ws in wb.worksheets:
                rows=[]
                for row in ws.iter_rows(values_only=True):
                    total+=1
                    if total>MAX_ROWS or len(row)>MAX_COLS: raise ValueError('Dimensions excessives')
                    rows.append([str(v).strip() if v is not None else '' for v in row])
                if rows: sheets.append({'name':ws.title,'headers':rows[0],'rows':rows[1:]})
        finally: wb.close()
    return {'sheets':sheets}
